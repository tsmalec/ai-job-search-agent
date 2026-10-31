"""ats_check.py — simulated screening report for the candidate's resume (say "simulated" when reporting).

Models how 2026 screening actually works:
  1. Parse: can a system read it? (pass/fail, not points)
  2. Required quals: each one met or not, like Workday HiredScore's LLM grading.
     All required met -> A/B; preferred quals split A from B.
  3. Title: recruiters search title + 2-3 skills. Target title appears as scope language
     (never as a title the candidate didn't hold).
  4. Keyword match % across every JD term (Jobscan-style). QC only, not a stop rule:
     enterprise ATS don't score keyword %; below 80% means look at the misses, not keep rebuilding.
  5. Readability: long bullets, soft openers, quantified bullets.
  6. Placement: each required qual should sit in a dated role, ideally the most recent;
     enterprise engines weight dated experience and the latest role over summary/skills lines.
  7. Acronyms (with --jd-text jd.txt): when the JD pairs a term with its acronym, the resume
     should carry both forms once. 8. Dates: one date format on every job line.
  9. Recruiter skim (eye-tracking research): the first two bullets of the latest role each
     carry a number; vague phrases flagged ("worked with stakeholders", "improved processes", ...).
Knockout questions are a separate check: scripts/knockouts.py jd.txt (step 2).
jd.json patterns are lowercased on load ("AI", "Power BI" now match) and short ones accept a
plural s ("okr" ~ "OKRs"); a degree qual in the Education section counts as placed; Acronyms only reads JD
pairs written "Spelled Out Form (ACR)", skips state codes, and accepts any "<words> (ACR)" already on the
resume; piping to head no longer crashes.

STOP when: parse PASS, required all met, title present. Keyword % is reported as a QC check.

Usage (never copy or edit this script): write a small jd.json with the JD's terms, then
       python3 scripts/ats_check.py RESUME.docx jd.json [--jd-text jd.txt]
  jd.json: {"title": "...", "title_core": [...], "required": {label: [variants]},
            "preferred": {...}, "keywords": {...}}   (see scripts/jd_example.json)
  label -> accepted honest variants; lowercase; a stem like "forecast" covers "forecasting".
  Without jd.json it runs the built-in sample terms below.
Use the JD's exact form at least once; the skills line is the fallback.
"""
import re, signal, sys, zipfile
from docx import Document

try:
    signal.signal(signal.SIGPIPE, signal.SIG_DFL)  # `| head` must not crash
except (AttributeError, ValueError):
    pass  # platform without this feature (e.g. Windows): nothing to do
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")  # ✓/⚠ on a Windows console
except Exception:
    pass  # not a reconfigurable stream

TITLE = "senior operations program manager"          # exact JD title
TITLE_CORE = ["program manager", "program management"]  # honest scope-language fallbacks
REQUIRED = {   # basic / minimum qualifications, one entry per requirement
    "program management": ["program management", "program manager"],
    "stakeholder management": ["stakeholder"],
    "process improvement": ["process improvement"],
}
PREFERRED = {  # preferred / nice-to-have qualifications
    "change management": ["change management"],
    "okr": ["okr"],
}
KEYWORDS = {   # other JD terms (responsibilities, tools, domain words) — skip pure filler
    "kpi": ["kpi"], "forecasting": ["forecast"], "governance": ["governance"],
    "cross-functional": ["cross-functional"], "dashboards": ["dashboard"], "ai": ["ai"],
}
VAGUE = ["worked with stakeholders", "worked closely with", "improved processes", "various", "assisted with",
         "helped with", "responsible for", "duties included", "tasked with", "involved in", "a variety of",
         "worked on", "in charge of", "handled various"]
SOFT_VERBS = {"served", "equipped", "helped", "assisted", "responsible", "worked", "participated"}
SECTIONS = ["SUMMARY", "EXPERIENCE", "EDUCATION", "SKILLS"]


def load(path):
    d = Document(path)
    paras = [p for p in d.paragraphs if p.text.strip()]
    text = "\n".join(p.text for p in paras)
    bullets = [p.text.strip() for p in paras
               if "List" in p.style.name or p._p.pPr is not None and p._p.pPr.numPr is not None]
    xml = zipfile.ZipFile(path).read("word/document.xml").decode("utf8", "ignore")
    header_text = " ".join(p.text for s in d.sections for hf in (s.header, s.footer)
                           for p in hf.paragraphs).strip()
    problems = []
    if d.tables: problems.append(f"{len(d.tables)} table(s)")
    if "txbxContent" in xml: problems.append("text box")
    if "<w:drawing" in xml or "<w:pict" in xml: problems.append("image/shape")
    if header_text: problems.append("text in header/footer")
    top = " ".join(p.text for p in paras[:3]).lower()
    if "@" not in top: problems.append("email not in first lines")
    missing_sec = [s for s in SECTIONS if s not in text.upper()]
    if missing_sec: problems.append("missing heading: " + ", ".join(missing_sec))
    regions, dates, first = split_regions(paras)
    fmts = {("full" if re.match(r"[A-Z][a-z]{3,}", d) else "abbr") for d in dates}
    bad = [d for d in dates if not DATE_RE.fullmatch(d)]
    if len(fmts) > 1 or bad: problems.append("mixed date formats: " + "; ".join(bad or dates))
    regions["first_bullets"] = first
    return text, bullets, problems, regions


DATE_RE = re.compile(r"[A-Z][a-z]+ \d{4} - (?:[A-Z][a-z]+ \d{4}|Present)")
MONTH = r"(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.? \d{4}"
JOBLINE = re.compile(r"\|\s*(" + MONTH + r"\s*[-–]\s*(?:" + MONTH + r"|Present))\s*$")


def split_regions(paras):
    """Text by region: summary, recent (first job), older (other jobs), skills, other. Plus job-line dates."""
    reg = {"summary": [], "recent": [], "older": [], "skills": [], "education": [], "other": []}
    dates, sec, job, first = [], None, -1, []
    for p in paras:
        t = p.text.strip()
        up = t.upper()
        if up in ("SUMMARY", "PROFESSIONAL SUMMARY"): sec = "summary"; continue
        if "EXPERIENCE" in up and len(t) < 30: sec = "exp"; continue
        if up.startswith("EDUCATION") and len(t) < 30: sec = "education"; continue
        if "SKILLS" in up and len(t) < 40: sec = "skills"; continue
        m = JOBLINE.search(t)
        if sec == "exp" and m:
            job += 1; dates.append(re.sub(r"\s*[-–]\s*", " - ", m.group(1)))
            continue
        if sec == "exp":
            reg["recent" if job <= 0 else "older"].append(t)
            if job == 0 and len(first) < 2 and "List" in p.style.name:
                first.append(t)
        elif sec in reg:
            reg[sec].append(t)
    return {k: low_norm(" ".join(v)) for k, v in reg.items()}, dates, first


def low_norm(text):
    low = " " + re.sub(r"[^a-z0-9$+&/ -]", " ", text.lower()) + " "
    return re.sub(r"\s+", " ", low)


ACR_SKIP = {"US", "USA", "EEO", "EOE", "LLC", "INC", "CEO", "CFO", "COO", "CTO", "OK", "AM", "PM", "ET", "PST", "EST"}
STATE_CODES = {"AL", "AK", "AZ", "AR", "CA", "CO", "CT", "DE", "FL", "GA", "HI", "ID", "IL", "IN", "IA", "KS", "KY", "LA",
               "ME", "MD", "MA", "MI", "MN", "MS", "MO", "MT", "NE", "NV", "NH", "NJ", "NM", "NY", "NC", "ND", "OH", "OK",
               "OR", "PA", "RI", "SC", "SD", "TN", "TX", "UT", "VT", "VA", "WA", "WV", "WI", "WY", "DC"}
SMALL = {"and", "of", "the", "for", "to", "in", "on", "&"}


def expansion(acr, text):
    """Spelled-out form of acr in text (initials match, small words skipped), or None."""
    words = re.findall(r"[A-Za-z][A-Za-z-]*|&", text)
    letters = acr.rstrip("s").upper()
    for i in range(len(words)):
        out, k, j = [], 0, i
        while j < len(words) and k < len(letters):
            w = words[j]
            if out and w.lower() in SMALL:
                out.append(w); j += 1; continue
            if w[0].upper() != letters[k]: break
            out.append(w); k += 1; j += 1
        if k == len(letters) and len(out) >= len(letters) >= 2 and out[0][0].isupper():
            return " ".join(out)
    return None


def paired(acr, text):
    """Spelled-out forms written right before "(ACR)" in text, checked against the initials."""
    out = []
    for m in re.finditer(r"((?:[A-Za-z][\w-]*\s+){1,7})\(" + re.escape(acr) + r"\)", text):
        exp = expansion(acr, m.group(1))
        if exp and m.group(1).rstrip().endswith(exp.split()[-1]):
            out.append(exp)
    return out


def acronym_report(jd_text, resume_text):
    lines = []
    for acr in sorted(set(re.findall(r"\b[A-Z]{2,6}s?\b", jd_text)) - ACR_SKIP - STATE_CODES):
        jd_pairs = paired(acr, jd_text)
        if not jd_pairs:
            continue  # only acronyms the JD itself spells out as "Form (ACR)"
        jd_exp = jd_pairs[0]
        r_acr = re.search(r"\b" + re.escape(acr.rstrip("s")) + r"s?\b", resume_text)
        r_exp = bool(paired(acr, resume_text)) or jd_exp.lower() in resume_text.lower()
        if r_acr and not r_exp:
            lines.append(f"{acr}: add '{jd_exp} ({acr})' once")
        elif r_exp and not r_acr:
            lines.append(f"{acr}: add the acronym after '{jd_exp}'")
    return lines


def hits(low, variants):
    """Patterns are matched lowercase; short ones (<=3 chars) as whole words with an optional plural s."""
    return any(re.search(r"\b" + re.escape(v.lower()) + r"s?\b", low) if len(v) <= 3 else v.lower() in low
               for v in variants)


def check(path, jd_text=None):
    text, bullets, problems, regions = load(path)
    low = " " + re.sub(r"[^a-z0-9$+&/ -]", " ", text.lower()) + " "
    low = re.sub(r"\s+", " ", low)

    req_ok = [k for k, v in REQUIRED.items() if hits(low, v)]
    req_miss = [k for k in REQUIRED if k not in req_ok]
    pref_ok = [k for k, v in PREFERRED.items() if hits(low, v)]
    pref_miss = [k for k in PREFERRED if k not in pref_ok]
    allkw = {**REQUIRED, **PREFERRED, **KEYWORDS}
    kw_ok = [k for k, v in allkw.items() if hits(low, v)]
    kw_miss = [k for k in allkw if k not in kw_ok]
    kw_pct = 100 * len(kw_ok) / max(1, len(allkw))

    title_exact = TITLE in low
    title_core = hits(low, TITLE_CORE)

    if not req_miss:
        grade = "A" if not PREFERRED or len(pref_ok) / len(PREFERRED) >= 0.5 else "B"
    else:
        grade = "C" if len(req_ok) >= len(REQUIRED) / 2 else "D"

    long_ = [b[:45] + "…" for b in bullets if len(b.split()) > 30]
    vague = sorted({v for b in bullets + [regions.get("summary", "")] for v in VAGUE if v in b.lower()})
    top = regions.get("first_bullets", [])
    top_nonum = [b[:45] + "…" for b in top if not re.search(r"\d", b)]
    soft = [b.split()[0] for b in bullets if b.split() and b.split()[0].lower() in SOFT_VERBS]
    quant = sum(bool(re.search(r"\d", b)) for b in bullets)

    ready = not problems and not req_miss and (title_exact or title_core)
    print("SIMULATED SCREEN (not a real ATS)")
    print(f"Parse: {'PASS' if not problems else 'FAIL — ' + '; '.join(problems)}")
    print(f"Est. grade: {grade}  |  Required {len(req_ok)}/{len(REQUIRED)}  |  "
          f"Preferred {len(pref_ok)}/{len(PREFERRED)}")
    if req_miss: print(f"  Required missing: {', '.join(req_miss)}")
    if pref_miss: print(f"  Preferred missing: {', '.join(pref_miss)}")
    print(f"Title: {'exact ✓' if title_exact else ('core words ✓ (exact missing)' if title_core else 'MISSING')}")
    weak = []
    for k in req_ok:
        v = REQUIRED[k]
        if hits(regions["recent"], v): continue
        if hits(regions.get("education", ""), v) and re.search(r"degree|bachelor|master|mba|b\.?s\b|ba\b|educat", k + " " + " ".join(v)):
            continue # a degree qual's home is the Education section
        weak.append(f"{k} (older role only)" if hits(regions["older"], v) else f"{k} (summary/skills only)")
    if weak: print(f"  Placement: move into a dated bullet, ideally the latest role: {', '.join(weak)}")
    print(f"Keywords {len(kw_ok)}/{len(allkw)} = {kw_pct:.0f}%  (QC: {'ok' if kw_pct >= 80 else 'under 80%, check the misses B honestly covers'})")
    if kw_miss: print(f"  Missing: {', '.join(kw_miss)}")
    if jd_text:
        acr = acronym_report(jd_text, text)
        print("Acronyms: " + ("both forms ✓" if not acr else "; ".join(acr)))
    print("Skim: first two bullets of the latest role " + ("both carry a number ✓" if top and not top_nonum
          else "need a number: " + "; ".join(top_nonum) if top else "not found"))
    if vague: print(f"  Vague phrases: {', '.join(vague)} (swap for the specific work in B)")
    print(f"Bullets {len(bullets)} | quantified {quant} | over 30 words {len(long_)} | soft openers {soft or 'none'}")
    print("STOP: ready to send" if ready else "KEEP GOING: see misses above")


if __name__ == "__main__":
    if any(a in ("-h", "--help") for a in sys.argv[1:]):  # never run on --help
        print((__doc__ or "").strip() or "No usage notes; see the script header.")
        sys.exit(0)
    jd_text = None
    if "--jd-text" in sys.argv:
        i = sys.argv.index("--jd-text")
        with open(sys.argv[i + 1], encoding="utf-8", errors="replace") as f:
            jd_text = f.read()
        del sys.argv[i:i + 2]
    if len(sys.argv) > 2:
        import json
        with open(sys.argv[2], encoding="utf-8") as f:
            jd = json.load(f)
        TITLE = jd["title"].lower()
        low_all = lambda d: {k: [x.lower() for x in v] for k, v in (d or {}).items()}
        TITLE_CORE = [x.lower() for x in jd.get("title_core", [])]
        REQUIRED, PREFERRED, KEYWORDS = low_all(jd.get("required")), low_all(jd.get("preferred")), low_all(jd.get("keywords"))
    check(sys.argv[1], jd_text)
