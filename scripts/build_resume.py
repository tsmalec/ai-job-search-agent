"""Build the candidate's resume .docx from a JSON content file. No template file needed.

Usage (from the repo root; never copy or edit this script):
    python3 scripts/build_resume.py content.json OUT.docx [--check]

content.json: see scripts/resume_example.json. Only write the content; the layout is fixed here.
--check renders a PDF with LibreOffice and prints the page count.

House layout (E_Style_ATS_Layout.md): single column, Calibri, no tables/text boxes/icons,
contact info as body text. Defaults are the one-page floor (10.5pt, 0.5" sides, 0.4" top/bottom);
override with a "layout" block in the JSON, e.g. {"font_pt": 11, "side_in": 0.7, "tb_in": 0.6}.

jobs: each item is {"title", "company", "dates", "bullets": [...]}.
Every role uses the same two lines (bold title, then "Company | dates"), bullets or not.
Education is two lines too (bold degree, then "School | year"). the candidate, Workday dropped the
one-line forms (bullet-less roles and the school), so never squash an entry onto one line.
"""
import json, subprocess, sys, tempfile
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, Inches

DEFAULTS = {"font": "Calibri", "font_pt": 10.5, "name_pt": 16, "heading_pt": 11,
            "side_in": 0.5, "tb_in": 0.4, "job_gap_pt": 6, "heading_gap_pt": 6}
sys.path.insert(0, str(Path(__file__).parent))
from common import CFG  # noqa: E402
EDU = CFG["candidate"]["education_line"]
CONTACT = CFG["candidate"]["contact_line"]


def build(c, out_path):
    L = {**DEFAULTS, **c.get("layout", {})}
    d = Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Inches(L["side_in"])
        s.top_margin = s.bottom_margin = Inches(L["tb_in"])
    normal = d.styles["Normal"]
    normal.font.name = L["font"]
    normal.font.size = Pt(L["font_pt"])
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), L["font"])
    pf = normal.paragraph_format
    pf.space_before = Pt(0); pf.space_after = Pt(0); pf.line_spacing = 1.0
    bullet = d.styles["List Bullet"]
    bullet.font.name = L["font"]; bullet.font.size = Pt(L["font_pt"])
    bullet.paragraph_format.space_after = Pt(0)

    def para(text="", bold=False, size=None, style=None, before=0, after=0, keep=False, center=False):
        p = d.add_paragraph(style=style)
        if text:
            r = p.add_run(text); r.bold = bold
            if size: r.font.size = Pt(size)
        f = p.paragraph_format
        f.space_before = Pt(before); f.space_after = Pt(after)
        if keep: f.keep_with_next = True
        if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return p

    def heading(text):
        para(text, bold=True, size=L["heading_pt"], before=L["heading_gap_pt"], after=2, keep=True)

    para(c.get("name", CFG["candidate"]["name"].upper()), bold=True, size=L["name_pt"])
    para(c.get("contact", CONTACT))
    heading("SUMMARY"); para(c["summary"])
    heading("WORK EXPERIENCE")
    for i, j in enumerate(c["jobs"]):
        gap = L["job_gap_pt"] if i else 0
        bullets = j.get("bullets") or []
        para(j["title"], bold=True, before=gap, keep=True)
        para(f"{j['company']} | {j['dates']}", keep=True)
        for b in bullets:
            para(b, style="List Bullet")
    heading("EDUCATION")
    degree, _, school = c.get("education", EDU).partition(" | ")
    para(degree, bold=True, keep=bool(school))
    if school:
        para(school)
    for sec in c.get("extra_sections", []):  # e.g. {"heading": "AI PROJECTS", "text": "..."}
        heading(sec["heading"]); para(sec["text"])
    heading(c.get("skills_heading", "SKILLS & CERTIFICATIONS")); para(c["skills"])
    d.save(out_path)


def _soffice():
    import shutil
    return shutil.which("soffice") or next(
        (p for p in [r"C:\Program Files\LibreOffice\program\soffice.exe",
                     r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
                     "/Applications/LibreOffice.app/Contents/MacOS/soffice"] if Path(p).exists()), None)


def page_count(docx_path):
    """Render to PDF with LibreOffice and count pages. Works on Linux and Windows (no pdfinfo needed)."""
    out = tempfile.mkdtemp()
    soffice = _soffice()
    if not soffice:
        print("⚠ LibreOffice not found: page count skipped (install it for --check page counts).", file=sys.stderr)
        return "?", Path(out) / "none.pdf"
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", out, str(docx_path)],
                   capture_output=True, timeout=180)
    pdf = Path(out) / (Path(docx_path).stem + ".pdf")
    if not pdf.exists():
        return "?", pdf
    data = pdf.read_bytes()
    import re as _re
    m = _re.findall(rb"/Type\s*/Pages\b[^>]*?/Count\s+(\d+)", data) or \
        _re.findall(rb"/Count\s+(\d+)[^>]*?/Type\s*/Pages\b", data)
    n = max(int(x) for x in m) if m else len(_re.findall(rb"/Type\s*/Page\b", data))
    return str(n), pdf


if __name__ == "__main__":
    if any(a in ("-h", "--help") for a in sys.argv[1:]):  # never run on --help
        print((__doc__ or "").strip() or "No usage notes; see the script header.")
        sys.exit(0)
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    content = json.loads(Path(args[0]).read_text())
    build(content, args[1])
    print(f"Built {args[1]}")
    if "--check" in sys.argv:
        pages, pdf = page_count(args[1])
        print(f"Pages: {pages}  (PDF: {pdf})")
